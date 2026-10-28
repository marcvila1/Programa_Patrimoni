#!/usr/bin/env python3
"""
Conversió de les fitxes Excel del Patrimoni de la DGPAMN a JSON normalitzat.

Ús:  python3 -I convert.py <carpeta_xlsx> <carpeta_sortida>

Entrada : llistes_<parc>.xlsx  (format transposat: col. A = camp, cols. B.. = elements)
          <sortida>/unitats_promotores.json  (còpia de la del repositori Motor-Contractacio-DGPA)
          scripts/mapa_unitats.json, scripts/correccions_manuals.json
Sortida : elements.json, espais.json, documents.json, catalegs.json, qualitat.json

Decisions de negoci aplicades:
  * Les còpies dels elements dels Ports que apareixen a llistes_secans.xlsx es descarten
    (les dades vàlides són les del fitxer dels Ports).
  * Un mateix element present a dos espais (casa_parcs_pirineus) es fusiona en un de sol,
    amb la llista `unitats`.
  * L'identificador d'element és <codi_unitat>/<codi> (p. ex. PNAE/acampada_cortalet);
    el codi sol no és únic. Els codis d'unitat són els del Motor de Contractació.
  * Es normalitzen valors (Sí/Si, variants ortogràfiques, camps compostos descomposats).
"""
import collections
import datetime
import glob
import json
import os
import re
import sys

import openpyxl

SRC = sys.argv[1]
OUT = sys.argv[2]
TODAY = datetime.date.today().isoformat()

avisos = []                       # coses que cal revisar a ull
corr = collections.Counter()      # recompte de correccions automàtiques


def avis(eid, camp, original, accio, estat=None):
    a = {"element": eid, "camp": camp, "original": original, "accio": accio}
    if estat:
        a["estat"] = estat
    avisos.append(a)


EXCEL_ERROR = re.compile(r"^#(REF!|N/A|VALUE!|DIV/0!|NAME\?|NULL!|NUM!)$")


# ----------------------------------------------------------------------------
# Helpers de neteja
# ----------------------------------------------------------------------------
def clean(v):
    """Text net: sense espais sobrers. Retorna None si és buit."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return v
    t = re.sub(r"\s+", " ", str(v)).strip()
    if t != str(v):
        corr["espais_sobrers"] += 1
    if EXCEL_ERROR.match(t):  # error de fórmula de l'Excel: no és una dada
        corr["error_formula_excel->null"] += 1
        return None
    return t or None


def sino(v, eid, camp, extra=None):
    t = clean(v)
    if t is None:
        avis(eid, camp, v, "valor buit; deixat a null")
        return None
    k = str(t).lower()
    if k in ("sí", "si", "sï"):
        if k != "sí":
            corr["Si->Sí"] += 1
        return True
    if k == "no":
        return False
    if extra and k in extra:
        return extra[k]
    avis(eid, camp, v, "valor Sí/No no reconegut; deixat a null")
    return None


def enum(v, mapping, eid, camp):
    t = clean(v)
    if t is None:
        return None
    k = str(t).lower()
    if k in mapping:
        return mapping[k]
    avis(eid, camp, v, "valor fora de catàleg; deixat a null")
    return None


def presencia(v, eid, camp):
    """Tipus de presència de personal -> (codi del catàleg, detall en text lliure).
    'Ocasional per arxivar documents' -> ('ocasional', 'per arxivar documents')."""
    t = clean(v)
    if t is None:
        return None, None
    k = str(t).lower()
    if k in MAP_PRES:
        return MAP_PRES[k], None
    if k == "sense personal":
        corr["presencia_personal_unificada"] += 1
        return "sense_presencia", None
    for paraula in ("permanent", "temporal", "estacional", "ocasional"):
        if k.startswith(paraula + " "):
            corr["presencia_amb_detall"] += 1
            return paraula, str(t)[len(paraula):].strip()
    avis(eid, camp, v, "valor fora de catàleg; deixat a null")
    return None, str(t)


def nota(v):
    t = clean(v)
    if t is None or str(t).lower() in ("no", "no."):
        return None
    return t


def sense_incidencia(v, prefix):
    """'No es detecten ...' -> None, resta de text es manté."""
    t = clean(v)
    if t is None or str(t).lower().startswith(prefix):
        return None
    return t


def int_or_none(v):
    try:
        return int(v)
    except (TypeError, ValueError):
        return None


def num_text(t):
    """Interpreta '1.213', '1,5', '15' -> int/float."""
    if "," in t:
        t = t.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(?:\.\d{3})+", t):
        t = t.replace(".", "")
    x = float(t)
    return int(x) if x == int(x) else x


# ----------------------------------------------------------------------------
# Catàlegs tancats (codi -> etiqueta)
# ----------------------------------------------------------------------------
CAT = {
    "tipus_espai": {
        "parc_natural": "Parc Natural", "parc_nacional": "Parc Nacional",
        "paratge_natural": "Paratge Natural d'Interès Nacional", "reserva_natural": "Reserva Natural",
    },
    "tipus_propietat": {"publica": "Pública", "privada": "Privada"},
    "tipus_finca": {"rustica": "Rústica", "urbana": "Urbana"},
    "tipus_acces": {
        "carretera_asfaltada": "Carretera asfaltada",
        "pista_asfaltada": "Pista per a vehicles asfaltada",
        "pista_sense_asfaltar": "Pista per a vehicles sense asfaltar",
        "pista_4x4": "Pista per a vehicles 4x4",
        "cami_terra": "Camí de terra",
        "cami_peatonal_pavimentat": "Camí peatonal pavimentat",
        "carrer_pavimentat": "Carrer pavimentat",
        "carrer_asfaltat": "Carrer asfaltat",
    },
    "sanejament": {"cap": "Sense sanejament", "fossa_septica": "Fossa sèptica",
                   "clavegueram_public": "Connexió al clavegueram públic"},
    "vigilancia": {"no": "Sense vigilància", "videocameres": "Videocàmeres"},
    "interior_adaptat": {"si": "Sí", "no": "No", "nomes_zones_publiques": "Només zones públiques"},
    "estat_conservacio": {"bo": "Bo", "regular": "Regular", "dolent": "Dolent",
                          "parcialment_ruinos": "Parcialment ruïnós", "ruinos": "Ruïnós"},
    "tipus_presencia_personal": {"sense_presencia": "Sense presència de personal", "permanent": "Permanent",
                                 "temporal": "Temporal", "estacional": "Estacional", "ocasional": "Ocasional"},
    "necessitat_permis_previ": {"no": "No", "si": "Sí", "en_desus": "En desús",
                                "autoritzacio_previa": "Autorització prèvia"},
    "estat_document": {"vigent": "Vigent", "caducat": "Caducat", "en_tramit": "En tràmit"},
    "precisio_any": {"exacte": "Any exacte", "decada": "Dècada", "aproximat": "Aproximat (text)",
                     "desconegut": "Desconegut"},
    "estat_reformes": {"desconegut": "Desconegut", "sense_reformes": "Sense reformes",
                       "amb_reformes": "Amb reformes"},
    "estat_dada_public": {"sense_dades": "Sense dades", "no_obert": "No obert al públic", "dada": "Amb dada"},
    "estat_pla": {"si": "Sí", "no": "No", "no_consta": "No consta"},
    "estat_execucio": {"planificacio": "Planificació", "operatiu": "Operatiu",
                       "requereix_revisio": "Requereix revisió", "suprimit": "Suprimit"},
    "fase_planificacio": {"f1_concepte": "Fase 1: Concepte", "f2_terrenys": "Fase 2: Terrenys",
                          "f3_avantprojecte": "Fase 3: Avantprojecte", "f4_projecte": "Fase 4: Projecte",
                          "f5_obres": "Fase 5: Obres"},
}


def rev(m):
    return {v.lower(): k for k, v in m.items()}


MAP_ACCES = {
    "carretera asfaltada": "carretera_asfaltada",
    "pista per vehicles asfaltada": "pista_asfaltada", "pista per a vehicles asfaltada": "pista_asfaltada",
    "pista per vehicles sense asfaltar": "pista_sense_asfaltar",
    "pista per a vehicles sense asfaltar": "pista_sense_asfaltar",
    "pista per vehicles 4x4": "pista_4x4", "pista per a vehicles 4x4": "pista_4x4",
    "camí de terra": "cami_terra",
    "camí peatonal pavimentat": "cami_peatonal_pavimentat",
    "camí pavimentat peatonal": "cami_peatonal_pavimentat",
    "carrer pavimentat": "carrer_pavimentat", "carrer asfaltat": "carrer_asfaltat",
}
MAP_SANEJ = {"no": "cap", "fossa sèptica": "fossa_septica",
             "connexió a clavagueram públic": "clavegueram_public",
             "connexió al clavagueram públic": "clavegueram_public",
             "connexió a xarxa de clavagueram": "clavegueram_public"}
MAP_VIG = {"no": "no", "videocàmeres": "videocameres"}
MAP_INT_ADAPT = {"sí": "si", "si": "si", "no": "no", "només zones públiques": "nomes_zones_publiques"}
MAP_ESTAT = rev(CAT["estat_conservacio"])
MAP_PRES = rev(CAT["tipus_presencia_personal"])
MAP_PERMIS = {"no": "no", "sí": "si", "si": "si", "en desús": "en_desus",
              "autorització prèvia": "autoritzacio_previa"}
MAP_PROP = {"pública": "publica", "privada": "privada"}
MAP_FINCA = {"rústica": "rustica", "urbana": "urbana"}

# Correccions ortogràfiques / unificació de redaccions de text lliure
TIPUS_ELEMENT_FIX = {
    "Casa Forestal": "Casa forestal",
    "Tallers i magatzem": "Tallers i magatzems",
    "Àrea interpretativa al aire lliure": "Àrea interpretativa a l'aire lliure",
}
REGIM_PERSONAL_FIX = {
    "sucontractat": "Subcontractat",
    "sucontractat o extern autoritzat": "Subcontractat o extern autoritzat",
    "personal de l'ajuntament de torroella": "Personal de l'Ajuntament de Torroella de Montgrí",
    "sense personal aliè": None, "sense personal del aliè": None,
}
REGIM_JURIDIC_FIX = {
    "acord col·laboració": "Acord de col·laboració",
    "cessió d'ús": "Cessió de l'ús",
    "cessió de domini": "Cessió del domini",
    "concessió d'ocupació": "Concessió de l'ocupació",
    "lloguer de l'ús": "Arrendament de l'ús",
}
PAVIMENT_FIX = {"asfalt": "Asfaltat"}


def fix_doc_text(t):
    t = re.sub(r"\bCoveni\b", "Conveni", t)
    t = re.sub(r"\bConvenió\b", "Conveni", t)
    t = t.replace("cessió d'ús", "cessió de l'ús").replace("cessió de domini", "cessió del domini")
    t = t.replace("la Entitat", "l'Entitat")
    t = t.replace("Concessió d'ocupació", "Concessió de l'ocupació")
    return t


# ----------------------------------------------------------------------------
# Parsers de camps compostos
# ----------------------------------------------------------------------------
def parse_utm(v, eid):
    t = clean(v)
    if t is None:
        return {"x": None, "y": None, "text_original": None}
    m = re.fullmatch(r"X0?(\d{6}),\s*Y(\d{7})", t)
    if m:
        return {"x": int(m.group(1)), "y": int(m.group(2)), "text_original": None}
    m = re.fullmatch(r"C?X0*(\d{6}),\s*[XY](\d{7})", t)
    if m:
        corr["utm_amb_typo_corregit"] += 1
        avis(eid, "coordenades_utm", t, f"typo evident; interpretat X={int(m.group(1))}, Y={int(m.group(2))}")
        return {"x": int(m.group(1)), "y": int(m.group(2)), "text_original": t}
    avis(eid, "coordenades_utm", t, "format no interpretable; x/y a null, text original conservat")
    return {"x": None, "y": None, "text_original": t}


def parse_altitud(v, eid):
    t = clean(v)
    if t is None:
        return None
    m = re.fullmatch(r"(\d[\d.]*)\s*m\.?", str(t))
    if not m:
        avis(eid, "altitud", t, "format no interpretable; deixat a null")
        return None
    return int(m.group(1).replace(".", "")) if re.fullmatch(r"\d{1,3}(?:\.\d{3})+", m.group(1)) \
        else int(m.group(1))


def parse_superficie(v, eid):
    t = clean(v)
    if t is None:
        return None
    m = re.fullmatch(r"([\d.,]+)\s*([mn])2", str(t))
    if not m:
        avis(eid, "superficie_m2", t, "format no interpretable; deixat a null")
        return None
    if m.group(2) == "n":
        corr["superficie_unitat_n2->m2"] += 1
        avis(eid, "superficie_m2", t, "unitat 'n2' interpretada com 'm2'")
    return num_text(m.group(1))


def parse_cp_municipi(v, eid):
    t = clean(v)
    m = re.fullmatch(r"CP(\d{5}) - (.+)", str(t or ""))
    if not m:
        avis(eid, "cp_i_municipi", t, "format no interpretable")
        return None, t
    return m.group(1), m.group(2).strip()


def parse_cadastre(v):
    t = clean(v)
    if t is None:
        return {"text": None, "poligon": None, "parcela": None}
    m = re.search(r"Pol[ií]gon[o]?(?:\s+cadastral)?\s+(\d+)\W+Parcel[·l]*a?\s+(\d+)", str(t), re.I)
    return {"text": t, "poligon": int(m.group(1)) if m else None, "parcela": int(m.group(2)) if m else None}


def parse_any(v, eid):
    if v is None:
        return {"valor": None, "text": None, "precisio": "desconegut"}
    if isinstance(v, int) or re.fullmatch(r"\d{4}", str(v).strip()):
        return {"valor": int(v), "text": None, "precisio": "exacte"}
    t = clean(v)
    if t.lower() == "desconegut":
        return {"valor": None, "text": None, "precisio": "desconegut"}
    m = re.fullmatch(r"Anys (\d0)", t)
    if m:
        d = int(m.group(1))
        return {"valor": (1900 if d >= 20 else 2000) + d, "text": t, "precisio": "decada"}
    return {"valor": None, "text": t, "precisio": "aproximat"}


def parse_reformes(v):
    t = clean(v)
    if t is None or str(t).lower() == "desconegut":
        return {"estat": "desconegut", "anys": [], "descripcio": None}
    if str(t).lower() == "sense reformes":
        return {"estat": "sense_reformes", "anys": [], "descripcio": None}
    anys = [int(a) for a in re.findall(r"\b(1[5-9]\d\d|20\d\d)\b", str(t))]
    return {"estat": "amb_reformes", "anys": anys, "descripcio": t}


def parse_reformes_prog(v):
    t = clean(v)
    if t is None or str(t).lower() == "no":
        return {"programades": False, "any": None, "descripcio": None}
    m = re.search(r"\b(20\d\d)\b", str(t))
    return {"programades": True, "any": int(m.group(1)) if m else None, "descripcio": t}


def parse_plantes(v, eid):
    t = clean(v)
    m = re.match(r"(\d+) plant", str(t or ""))
    if not m:
        avis(eid, "nombre_plantes", t, "format no interpretable")
        return {"plantes": None, "altell": False}
    return {"plantes": int(m.group(1)), "altell": "altell" in str(t).lower()}


def material(v, prefixos_nul):
    t = clean(v)
    if t is None:
        return None
    k = str(t).lower()
    if any(k.startswith(p) for p in prefixos_nul):
        return None
    if k in PAVIMENT_FIX:
        corr["paviment_unificat"] += 1
        return PAVIMENT_FIX[k]
    return t


def doc_fitxer(v):
    t = clean(v)
    if t is None or str(t).lower() in ("no", "no consta"):
        return None
    return t if str(t).lower().endswith(".pdf") else None


def parse_volum(v):
    t = clean(v)
    if t is None or str(t).lower() == "sense dades":
        return {"estat": "sense_dades", "valor": None}
    if str(t).lower().startswith("no obert"):
        return {"estat": "no_obert", "valor": None}
    m = re.fullmatch(r"([\d.]+) persones", str(t))
    if m:
        return {"estat": "dada", "valor": int(m.group(1).replace(".", ""))}
    return {"estat": "dada", "valor": None, "text": t}


def parse_perfil(v):
    t = clean(v)
    if t is None or str(t).lower() == "sense dades":
        return {"estat": "sense_dades", "text": None}
    if str(t).lower().startswith("no obert"):
        return {"estat": "no_obert", "text": None}
    return {"estat": "dada", "text": t}


def parse_pla(v):
    t = clean(v)
    if t is None or str(t).lower() == "no":
        return {"estat": "no", "fitxer": None}
    if str(t).lower() == "no consta":
        return {"estat": "no_consta", "fitxer": None}
    return {"estat": "si", "fitxer": t}


def parse_doc_juridic(tipus_v, fitxer_v):
    t = clean(tipus_v)
    estat = None
    if t is not None and str(t).lower() in ("no", "no consta"):
        t = None
    if t is not None:
        t = fix_doc_text(t)
        low = t.lower()
        if "caducat" in low:
            estat = "caducat"
        elif "tràmit" in low or "tramitació" in low:
            estat = "en_tramit"
        t = re.sub(r"\s*\((?:caducat|caducada|en tràmit|en tramitació)\)", "", t).strip()
    return {"tipus": t, "estat": estat, "fitxer": doc_fitxer(fitxer_v), "url": None}


def nullish(v, extra=()):
    t = clean(v)
    if t is None or str(t).lower() in ("no consta", "desconegut", "sense dades") + tuple(extra):
        return None
    return t


# ----------------------------------------------------------------------------
# Lectura dels Excel
# ----------------------------------------------------------------------------
files = sorted(glob.glob(os.path.join(SRC, "llistes_*.xlsx")))
if not files:
    sys.exit("No s'han trobat fitxers llistes_*.xlsx a " + SRC)

NROWS = 154
raw = []
names = None
for f in files:
    slug = re.sub(r"^llistes_|\.xlsx$", "", os.path.basename(f))
    ws = openpyxl.load_workbook(f, data_only=True).active
    if names is None:
        names = [ws.cell(r, 1).value for r in range(1, NROWS + 1)]
    for c in range(2, ws.max_column + 1):
        if ws.cell(1, c).value is None:
            continue
        v = {r: ws.cell(r, c).value for r in range(1, NROWS + 1)}
        raw.append({"slug": slug, "fitxer": os.path.basename(f),
                    "col": openpyxl.utils.get_column_letter(c), "v": v})

# Parc "propi" de cada fitxer. Un fitxer amb un sol parc n'és el propietari; en un fitxer
# mixt (Secans = Secans + còpies dels Ports) el parc propi és l'únic que no té fitxer propi.
slugs = sorted({e["slug"] for e in raw})
parcs_per_fitxer = {s: collections.Counter(e["v"][3] for e in raw if e["slug"] == s) for s in slugs}
owner = {}
for s in slugs:
    if len(parcs_per_fitxer[s]) == 1:
        owner[next(iter(parcs_per_fitxer[s]))] = s
file_parc = {}
for s in slugs:
    if len(parcs_per_fitxer[s]) == 1:
        file_parc[s] = next(iter(parcs_per_fitxer[s]))
    else:
        propis = [n for n in parcs_per_fitxer[s] if n not in owner]
        if len(propis) != 1:
            sys.exit(f"No es pot determinar el parc propi de {s}: {dict(parcs_per_fitxer[s])}")
        file_parc[s] = propis[0]
        owner[propis[0]] = s

# Unitats promotores (font de veritat dels codis i noms) i mapatge fitxer -> unitat
_dir = os.path.dirname(os.path.abspath(__file__))
with open(os.path.join(_dir, "mapa_unitats.json"), encoding="utf-8") as fh:
    MAPA = {k: v for k, v in json.load(fh).items() if not k.startswith("_")}
unitats_path = os.path.join(OUT, "unitats_promotores.json")
if not os.path.exists(unitats_path):
    sys.exit("Falta " + unitats_path + " (còpia de unitats_promotores.json del repositori del Motor)")
with open(unitats_path, encoding="utf-8") as fh:
    _unitats_json = json.load(fh)
UNITATS = {u["codi"]: u for u in _unitats_json["unitats"]}
UNITATS_META = _unitats_json.get("metadata", {})
for _slug in {e["slug"] for e in raw}:
    if _slug not in MAPA:
        sys.exit(f"Fitxer sense unitat assignada a mapa_unitats.json: {_slug}")
    if MAPA[_slug] not in UNITATS:
        sys.exit(f"La unitat {MAPA[_slug]} no existeix a unitats_promotores.json")

# 1) Descartar còpies d'elements d'un parc que viuen al fitxer d'un altre (Ports a Secans)
kept, descartats = [], []
for e in raw:
    if e["v"][3] != file_parc[e["slug"]]:
        os_ = owner[e["v"][3]]
        ref = next((x for x in raw if x["slug"] == os_ and x["v"][1] == e["v"][1]), None)
        if ref is None:
            sys.exit(f"Còpia sense original: {e['slug']}/{e['v'][1]}")
        dif = [{"camp": names[r - 1], "fila": r, "valor_descartat": e["v"][r], "valor_conservat": ref["v"][r]}
               for r in range(1, NROWS + 1) if e["v"][r] != ref["v"][r]]
        descartats.append({"codi": e["v"][1], "fitxer": e["fitxer"], "columna": e["col"],
                           "unitat_propietaria": MAPA[os_], "diferencies": dif})
        continue
    kept.append(e)

# 2) Fusionar elements compartits entre parcs; separar col·lisions de codi
by_codi = collections.defaultdict(list)
for e in kept:
    by_codi[e["v"][1]].append(e)

elements_raw, compartits = [], []
for codi, grup in by_codi.items():
    if len(grup) == 1:
        elements_raw.append((grup[0], [grup[0]["slug"]]))
        continue
    claus = {(g["v"][5], g["v"][11], g["v"][17]) for g in grup}
    if len(claus) == 1:  # mateix nom, cadastre i adreça => mateix element
        base = grup[0]
        dif = [{"camp": names[r - 1], "fila": r, "valors": {MAPA[g["slug"]]: g["v"][r] for g in grup}}
               for r in range(1, NROWS + 1) if len({str(g["v"][r]) for g in grup}) > 1]
        compartits.append({"codi": codi, "unitats": [MAPA[g["slug"]] for g in grup],
                           "base_conservada": MAPA[base["slug"]], "diferencies": dif})
        elements_raw.append((base, [g["slug"] for g in grup]))
    else:  # mateix codi, elements diferents
        for g in grup:
            elements_raw.append((g, [g["slug"]]))
            avis(f"{MAPA[g['slug']]}/{codi}", "codi", codi,
                 "codi repetit en una altra unitat amb un element diferent; l'ID inclou la unitat")

park_order = sorted({e["slug"] for e in raw})
elements_raw.sort(key=lambda t: (park_order.index(t[0]["slug"]), openpyxl.utils.column_index_from_string(t[0]["col"])))

# ----------------------------------------------------------------------------
# Construcció dels elements
# ----------------------------------------------------------------------------
ZONA_BASE = 71
ZONA_N = 7
ZONA_STRIDE = 12


def construir_zona(v, n, eid):
    b = ZONA_BASE + ZONA_STRIDE * (n - 1)
    tp, tp_detall = presencia(v[b + 9], eid, f"zona{n}.tipus_presencia_personal")
    return {
        "numero": n,
        "nom": clean(v[b]),
        "usos": clean(v[b + 1]),
        "usuaris": clean(v[b + 2]),
        "accessible_public": sino(v[b + 3], eid, f"zona{n}.accessible_public"),
        "accessible_menors": sino(v[b + 4], eid, f"zona{n}.accessible_menors"),
        "adaptat_discapacitats": sino(v[b + 5], eid, f"zona{n}.adaptat_discapacitats"),
        "presencia_personal_parc": sino(v[b + 6], eid, f"zona{n}.presencia_personal_parc"),
        "presencia_personal_alie": sino(v[b + 7], eid, f"zona{n}.presencia_personal_alie"),
        "regim_personal_alie": regim_personal(v[b + 8]),
        "tipus_presencia_personal": tp,
        "detall_presencia_personal": tp_detall,
        "necessitat_permis_previ": enum(v[b + 10], MAP_PERMIS, eid, f"zona{n}.necessitat_permis_previ"),
        "inclos_pla_autoproteccio": sino(v[b + 11], eid, f"zona{n}.inclos_pla_autoproteccio"),
    }


def regim_personal(v):
    t = clean(v)
    if t is None:
        return None
    k = str(t).lower()
    if k in REGIM_PERSONAL_FIX:
        nou = REGIM_PERSONAL_FIX[k]
        if nou is None:
            corr["sense_personal_alie->null"] += 1
        elif nou != t:
            corr["regim_personal_alie_unificat"] += 1
        return nou
    return t


elements = []
for e, slugs_el in elements_raw:
    parcs = [MAPA[s] for s in slugs_el]  # codis d'unitat
    v = e["v"]
    codi = clean(v[1])
    eid = f"{parcs[0]}/{codi}"

    # --- identificació
    tipus_el = clean(v[7])
    if tipus_el in TIPUS_ELEMENT_FIX:
        corr["tipus_element_unificat"] += 1
        tipus_el = TIPUS_ELEMENT_FIX[tipus_el]
    altres = clean(v[6])
    altres_noms = [] if altres is None or str(altres).lower() == "no" else \
        [x.strip() for x in re.split(r",\s*", str(altres)) if x.strip()]
    refcat = clean(v[11])
    refcat_nota = None
    if refcat is not None and not re.fullmatch(r"[0-9A-Z]{14,20}", str(refcat)):
        refcat_nota = refcat
        avis(eid, "referencia_cadastral", refcat, "no és una referència cadastral; passada a nota")
        refcat = None
    nom = clean(v[5])
    if nom is None and clean(v[2]):
        nom = clean(v[2])
        if v[5] is not None and EXCEL_ERROR.match(str(v[5]).strip()):
            corr["nom_error_formula->titol"] += 1
            avis(eid, "nom", v[5], "error de fórmula de l'Excel al camp Nom; s'ha pres el Títol", estat="corregit_automatic")
    identificacio = {
        "nom": nom, "titol": clean(v[2]), "altres_noms": altres_noms,
        "tipus_element": tipus_el, "destinacio_principal": clean(v[8]),
        "tipus_propietat": enum(v[9], MAP_PROP, eid, "tipus_propietat"),
        "titular": nullish(v[10]),
        "referencia_cadastral": refcat, "referencia_cadastral_nota": refcat_nota,
        "finca_registral": nullish(v[12]),
        "tipus_finca": enum(v[13], MAP_FINCA, eid, "tipus_finca"),
        "proteccio_patrimoni": sino(v[14], eid, "proteccio_patrimoni"),
    }

    # --- localització
    cp, municipi = parse_cp_municipi(v[18], eid)
    adreca = clean(v[17])
    localitzacio = {
        "adreca": None if adreca is None or str(adreca).lower() == "sense adreça" else adreca,
        "cp": cp, "municipi": municipi,
        "localitzacio_cadastre": parse_cadastre(v[19]),
        "tipus_acces": enum(v[20], MAP_ACCES, eid, "tipus_acces"),
        "acces_adaptat": sino(v[21], eid, "acces_adaptat"),
        "utm": parse_utm(v[22], eid),
        "altitud_m": parse_altitud(v[23], eid),
    }

    # --- construcció i estat
    obert = clean(v[42])
    if obert is not None and str(obert).lower() == "bi":
        # Typo; l'element no és accessible al públic segons la resta de camps (zones_accessibles = No)
        obert_val = False
        corr["Bi->No"] += 1
        avis(eid, "obert_public", obert, "typo 'Bi'; interpretat com a No perquè 'zones accessibles al públic' = No")
    else:
        obert_val = sino(v[42], eid, "obert_public")
    nombre_fotos_def = clean(v[49])
    construccio = {
        "any_construccio": parse_any(v[26], eid),
        "reformes": parse_reformes(v[27]),
        "reformes_programades": parse_reformes_prog(v[28]),
        **parse_plantes(v[29], eid),
        "superficie_m2": parse_superficie(v[30], eid),
        "planols": {"fitxer": doc_fitxer(v[31]), "url": None},
        "materials": {
            "envolvent": material(v[32], ("sense ",)),
            "coberta": material(v[33], ("sense ",)),
            "estructura": material(v[34], ("sense ",)),
            "paviments_interiors": material(v[35], ("no pavimentat", "sense pavimentar", "sense paviment")),
        },
        "serveis": {
            "electricitat": sino(v[36], eid, "electricitat"),
            "aigua": sino(v[37], eid, "aigua"),
            "sanejament": enum(v[38], MAP_SANEJ, eid, "sanejament"),
            "internet": sino(v[39], eid, "internet"),
            "alarma": sino(v[40], eid, "alarma"),
            "vigilancia": enum(v[41], MAP_VIG, eid, "vigilancia"),
        },
        "obert_public": obert_val,
        "interior_adaptat_discapacitats": enum(v[43], MAP_INT_ADAPT, eid, "interior_adaptat_discapacitats"),
        "estat_conservacio": enum(v[44], MAP_ESTAT, eid, "estat_conservacio"),
        "deficiencies": {
            "detectades": (False if str(clean(v[45]) or "").lower().startswith("no")
                           else sino(v[45], eid, "deficiencies_detectades")),
            "descripcio": sense_incidencia(v[47], "no es detecten"),
            "reparacions_i_cost": sense_incidencia(v[48], "no es detecten"),
            "nombre_fotografies": 0 if nombre_fotos_def is None or str(nombre_fotos_def).lower() == "no"
            else int_or_none(nombre_fotos_def),
        },
        "riscos": {
            "detectats": (False if str(clean(v[46]) or "").lower().startswith("no")
                          else sino(v[46], eid, "riscos_detectats")),
            "descripcio": sense_incidencia(v[50], "no es detecten"),
        },
    }

    # --- ús i règim jurídic
    regim = clean(v[59])
    if regim is not None:
        if str(regim).lower() in ("no consta", "desconegut"):
            regim = None
        elif str(regim).lower() in REGIM_JURIDIC_FIX:
            corr["regim_juridic_unificat"] += 1
            regim = REGIM_JURIDIC_FIX[str(regim).lower()]
    us_i_regim = {
        "us_principal": clean(v[56]),
        "propietari": nullish(v[57]),
        "usuari_explotador": nullish(v[58]),
        "regim_juridic": regim,
        "document_juridic": parse_doc_juridic(v[60], v[61]),
        "zones_accessibles_public": sino(v[62], eid, "zones_accessibles_public"),
        "presencia_personal_parc": sino(v[63], eid, "presencia_personal_parc"),
        "presencia_personal_alie": sino(v[64], eid, "presencia_personal_alie"),
        "tipus_presencia_personal": presencia(v[65], eid, "tipus_presencia_personal")[0],
        "detall_presencia_personal": presencia(v[65], eid, "tipus_presencia_personal")[1],
        "volum_anual_public": parse_volum(v[66]),
        "perfil_public": parse_perfil(v[67]),
        "pla_autoproteccio": parse_pla(v[68]),
    }

    # --- zones
    nz = int_or_none(v[70]) or 1
    zones = [construir_zona(v, n, eid) for n in range(1, min(nz, ZONA_N) + 1)]
    for n in range(nz + 1, ZONA_N + 1):  # zones més enllà del nombre declarat: han d'estar buides (0)
        b = ZONA_BASE + ZONA_STRIDE * (n - 1)
        if any(v[r] not in (0, None, "") for r in range(b, b + ZONA_STRIDE)):
            avis(eid, f"zona{n}", None, "dades a una zona per sobre de 'Nombre de zones'; no migrades")

    elements.append({
        "id": eid, "codi": codi, "unitat": parcs[0], "unitats": parcs,
        "identificacio": identificacio, "localitzacio": localitzacio, "construccio": construccio,
        "us_i_regim": us_i_regim,
        "estat_execucio": {"estat": "operatiu", "fase": None, "motius_automatics": [], "motiu": None},
        "zones": zones,
        "fotografies": {"nombre": int_or_none(v[53]) or 0, "fitxers": []},
        "notes": {"identificacio": nota(v[15]), "localitzacio": nota(v[24]), "construccio": nota(v[51]),
                  "fotografies": nota(v[54]), "us": nota(v[69])},
        "meta": {"origen": {"fitxer": e["fitxer"], "columna": e["col"]}, "creat": TODAY, "modificat": None,
                 "modificat_per": None, "versio": 1, "pendents": []},
    })

# ----------------------------------------------------------------------------
# Correccions manuals (scripts/correccions_manuals.json)
# ----------------------------------------------------------------------------
def set_path(obj, path, valor, eid):
    claus = path.split(".")
    for k in claus[:-1]:
        obj = obj[k]
    if claus[-1] not in obj:
        sys.exit(f"Camp inexistent en la correcció manual: {eid} {path}")
    obj[claus[-1]] = valor


for a in avisos:
    a.setdefault("estat", "informatiu" if a["camp"] == "codi" else "obert")

man_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), "correccions_manuals.json")
man = {}
if os.path.exists(man_path):
    with open(man_path, encoding="utf-8") as fh:
        man = json.load(fh)
    per_id = {el["id"]: el for el in elements}

    def estat_avis(element, camp_avis, estat, nota_):
        trobats = [a for a in avisos if a["element"] == element and a["camp"] == camp_avis]
        if not trobats:
            sys.exit(f"La decisió manual no correspon a cap avís: {element} {camp_avis}")
        for a in trobats:
            a["estat"] = estat
            a["decisio"] = nota_

    for c in man.get("correccions", []):
        if c["element"] not in per_id:
            sys.exit(f"Element inexistent: {c['element']}")
        set_path(per_id[c["element"]], c["camp"], c["valor"], c["element"])
        per_id[c["element"]]["meta"]["modificat"] = TODAY
        corr["correccio_manual"] += 1
        estat_avis(c["element"], c["camp_avis"], "corregit_manual", c["motiu"])
    for p in man.get("pendents", []):
        if p["element"] not in per_id:
            sys.exit(f"Element inexistent: {p['element']}")
        per_id[p["element"]]["meta"]["pendents"].append({"camp": p["camp"], "motiu": p["motiu"]})
        corr["pendent_de_dada"] += 1
        estat_avis(p["element"], p["camp_avis"], "pendent", p["motiu"])
    for c in man.get("confirmacions", []):
        estat_avis(c["element"], c["camp_avis"], "confirmat", c["decisio"])
        corr["avis_confirmat"] += 1
    for c in man.get("compartits", []):
        trobats = [x for x in compartits if x["codi"] == c["codi"]]
        if not trobats:
            sys.exit(f"Decisió sobre un element compartit inexistent: {c['codi']}")
        for x in trobats:
            x["decisio"] = c["decisio"]

# ----------------------------------------------------------------------------
# Estat d'execució. Tots els immobles inventariats són «operatiu» o «requereix_revisio».
# Regla automàtica (la mateixa que aplica index.html en desar):
#   - document jurídic caducat (el contracte ha deixat de ser vigent)
#   - contracte de paraula (acord verbal) al règim jurídic o al tipus de document
# ----------------------------------------------------------------------------
def motius_execucio(el):
    u = el["us_i_regim"]
    d = u.get("document_juridic") or {}
    motius = []
    if d.get("estat") == "caducat":
        doc = f"«{d['tipus']}» " if d.get("tipus") else ""
        motius.append({"codi": "contracte_no_vigent",
                       "text": f"El document jurídic {doc}ha deixat de ser vigent (caducat): cal renovar-lo o regularitzar l'ús."})
    if any("acord verbal" in (t or "").lower() for t in (u.get("regim_juridic"), d.get("tipus"))):
        motius.append({"codi": "contracte_paraula",
                       "text": "L'ús es basa en un contracte de paraula (acord verbal) sense document formal: cal formalitzar-lo."})
    return motius


# Conserva la data de creació original si l'element ja existia en una conversió anterior
prev_path = os.path.join(OUT, "elements.json")
if os.path.exists(prev_path):
    with open(prev_path, encoding="utf-8") as fh:
        prev_creat = {p["id"]: p.get("meta", {}).get("creat") for p in json.load(fh)}
    for el in elements:
        if prev_creat.get(el["id"]):
            el["meta"]["creat"] = prev_creat[el["id"]]

for el in elements:
    ex = el["estat_execucio"]
    ex["motius_automatics"] = motius_execucio(el)
    if ex["estat"] in ("operatiu", "requereix_revisio"):
        ex["estat"] = "requereix_revisio" if ex["motius_automatics"] else "operatiu"
    corr[f"estat_execucio_{ex['estat']}"] += 1

# ----------------------------------------------------------------------------
# Parcs, documents, catàlegs, qualitat
# ----------------------------------------------------------------------------
def tipus_espai(nom):
    n = nom.lower()
    if n.startswith("parc nacional"):
        return "parc_nacional"
    if n.startswith("parc natural"):
        return "parc_natural"
    if n.startswith("paratge"):
        return "paratge_natural"
    if n.startswith("reserva"):
        return "reserva_natural"
    return None


# espais.json: dades pròpies del Patrimoni per a cada unitat que és un espai natural.
# El nom, NIF i seu s'obtenen de unitats_promotores.json (no es dupliquen aquí).
espais = {}
noms_unitats = []
for slug in park_order:
    cod = MAPA[slug]
    nom_excel = file_parc[slug]
    if "Interés" in nom_excel:
        corr["nom_parc_Interés->Interès"] += 1
        nom_excel = nom_excel.replace("Interés", "Interès")
    u = UNITATS[cod]
    if u["tipus"] != "ESPAI_NATURAL":
        avis(cod, "unitat.tipus", u["tipus"], "la unitat associada a un fitxer no és ESPAI_NATURAL")
    espais[cod] = {
        "tipus_espai": tipus_espai(u["nom"]),
        "fitxer_origen": f"llistes_{slug}.xlsx",
        "nom_excel": nom_excel,
        "nombre_elements": sum(1 for el in elements if cod in el["unitats"]),
    }
    noms_unitats.append({"unitat": cod, "nom_excel": nom_excel, "nom_unitat": u["nom"],
                         "coincideix": nom_excel == u["nom"]})

docs = collections.defaultdict(lambda: {"fitxer": None, "url": None, "referenciat_per": []})
for el in elements:
    for tipus, f in (("planols", el["construccio"]["planols"]["fitxer"]),
                     ("document_juridic", el["us_i_regim"]["document_juridic"]["fitxer"]),
                     ("pla_autoproteccio", el["us_i_regim"]["pla_autoproteccio"]["fitxer"])):
        if f:
            d = docs[f]
            d["fitxer"] = f
            d["referenciat_per"].append({"element": el["id"], "camp": tipus})
documents = sorted(docs.values(), key=lambda d: d["fitxer"])

oberts = {
    "tipus_element": sorted({el["identificacio"]["tipus_element"] for el in elements if el["identificacio"]["tipus_element"]}),
    "municipis": sorted({el["localitzacio"]["municipi"] for el in elements if el["localitzacio"]["municipi"]}),
    "titulars": sorted({el["identificacio"]["titular"] for el in elements if el["identificacio"]["titular"]}),
    "propietaris": sorted({el["us_i_regim"]["propietari"] for el in elements if el["us_i_regim"]["propietari"]}),
    "usuaris_explotadors": sorted({el["us_i_regim"]["usuari_explotador"] for el in elements if el["us_i_regim"]["usuari_explotador"]}),
    "regims_juridics": sorted({el["us_i_regim"]["regim_juridic"] for el in elements if el["us_i_regim"]["regim_juridic"]}),
    "tipus_document_juridic": sorted({el["us_i_regim"]["document_juridic"]["tipus"] for el in elements if el["us_i_regim"]["document_juridic"]["tipus"]}),
    "regims_personal_alie": sorted({z["regim_personal_alie"] for el in elements for z in el["zones"] if z["regim_personal_alie"]}),
}
catalegs = {"tancats": CAT, "oberts": oberts}

qualitat = {
    "generat": TODAY,
    "resum": {
        "fitxers_xlsx": len(files), "columnes_llegides": len(raw),
        "copies_descartades_secans_ports": len(descartats),
        "elements_compartits_fusionats": len(compartits),
        "elements_resultants": len(elements), "unitats_amb_patrimoni": len(espais),
        "documents_referenciats": len(documents),
        "avisos_totals": len(avisos),
        "avisos_per_estat": dict(collections.Counter(a["estat"] for a in avisos)),
    },
    "decisions": man.get("decisions", []),
    "correccions_automatiques": dict(corr),
    "avisos": avisos,
    "elements_compartits": compartits,
    "copies_descartades": descartats,
    "noms_espais_excel_vs_unitat": noms_unitats,
}


def dump(nom, obj):
    with open(os.path.join(OUT, nom), "w", encoding="utf-8") as fh:
        json.dump(obj, fh, ensure_ascii=False, indent=2)
        fh.write("\n")


os.makedirs(OUT, exist_ok=True)
dump("espais.json", espais)
dump("elements.json", elements)
dump("documents.json", documents)
dump("catalegs.json", catalegs)
dump("qualitat.json", qualitat)
metadata = {
    "projecte": "Motor Patrimoni — DGPA",
    "generat": TODAY,
    "fonts": {
        "fitxes_excel": "marcvila1/Programa_Patrimoni (llistes_*.xlsx)",
        "unitats_promotores": {
            "repositori": "marcvila1/Motor-Contractacio-DGPA", "fitxer": "unitats_promotores.json",
            "versio": UNITATS_META.get("versio"), "data_actualitzacio": UNITATS_META.get("data_actualitzacio"),
        },
    },
    **man.get("metadades", {}),
    "recomptes": {"unitats_amb_patrimoni": len(espais), "elements": len(elements),
                  "documents": len(documents), "elements_amb_pendents": sum(1 for el in elements if el["meta"]["pendents"]),
                  "estat_execucio": dict(collections.Counter(el["estat_execucio"]["estat"] for el in elements))},
}
dump("metadata.json", metadata)
print(json.dumps(qualitat["resum"], ensure_ascii=False, indent=2))
print("Correccions:", dict(corr))
